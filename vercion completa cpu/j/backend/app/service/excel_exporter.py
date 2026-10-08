"""Genera el reporte Excel (.xlsx) de los cheques analizados."""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.enums import FieldKey
from app.exception import ExportException
from app.model import CheckRecord

_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_STATUS_FILL = {
    "COMPLETO": PatternFill("solid", fgColor="D9F2E3"),
    "REVISAR": PatternFill("solid", fgColor="FFF1CC"),
    "ERROR": PatternFill("solid", fgColor="F9D6D5"),
}
_BORDER = Border(*(Side(style="thin", color="D0D7E2"),) * 4)

COLUMNS = [
    ("ID", 7), ("Archivo", 28), ("Estado", 12), ("Banco", 26), ("N° Cheque", 14), ("Ciudad", 14), ("Fecha", 12),
    ("Beneficiario", 34), ("Monto", 14), ("Moneda", 9), ("Monto en letras", 48), ("Valor en letras", 14),
    ("¿Montos coinciden?", 12), ("Cuenta", 18), ("Código de ruta", 14), ("Línea MICR", 34), ("Concepto", 26),
    ("Firma", 8), ("Confianza", 11), ("Observaciones", 60), ("Procesado", 20), ("Tiempo (ms)", 11),
]


def _amount(value) -> Decimal | None:
    try:
        return Decimal(value) if value not in (None, "") else None
    except Exception:
        return None


def build_excel(records: list[CheckRecord]) -> bytes:
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Cheques"
        ws.append([name for name, _ in COLUMNS])
        for col, (_, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 30

        for r in records:
            match = {True: "Sí", False: "No"}.get(r.amount_match, "-")
            ws.append([
                r.id, r.file_name, r.status, r.value(FieldKey.BANCO), r.value(FieldKey.NUMERO_CHEQUE),
                r.value(FieldKey.CIUDAD),
                datetime.fromisoformat(r.value(FieldKey.FECHA)).date() if r.value(FieldKey.FECHA) else None,
                r.value(FieldKey.BENEFICIARIO), _amount(r.value(FieldKey.MONTO)), r.value(FieldKey.MONEDA),
                r.value(FieldKey.MONTO_LETRAS), r.amount_words_value, match, r.value(FieldKey.CUENTA),
                r.value(FieldKey.CODIGO_RUTA), r.value(FieldKey.LINEA_MICR), r.value(FieldKey.CONCEPTO),
                r.value(FieldKey.FIRMA), r.overall_confidence, " ".join(r.observations),
                r.processed_at.replace("T", " "), r.processing_ms,
            ])
            row = ws.max_row
            for col in range(1, len(COLUMNS) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = _BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=col in (8, 11, 20))
            ws.cell(row=row, column=3).fill = _STATUS_FILL.get(r.status, PatternFill())
            ws.cell(row=row, column=7).number_format = "DD/MM/YYYY"
            ws.cell(row=row, column=9).number_format = "#,##0.00"
            ws.cell(row=row, column=12).number_format = "#,##0.00"
            ws.cell(row=row, column=19).number_format = "0%"

        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        _summary_sheet(wb, records)
        _ocr_sheet(wb, records)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
    except Exception as exc:
        raise ExportException(str(exc)) from exc


def _summary_sheet(wb: Workbook, records: list[CheckRecord]) -> None:
    ws = wb.create_sheet("Resumen")
    total = sum((_amount(r.value(FieldKey.MONTO)) or Decimal(0)) for r in records)
    rows = [
        ("Reporte generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Cheques analizados", len(records)),
        ("Completos", sum(r.status == "COMPLETO" for r in records)),
        ("Por revisar", sum(r.status == "REVISAR" for r in records)),
        ("Con error", sum(r.status == "ERROR" for r in records)),
        ("Monto total", total),
    ]
    for label, value in rows:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=6, column=2).number_format = "#,##0.00"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 22


def _ocr_sheet(wb: Workbook, records: list[CheckRecord]) -> None:
    ws = wb.create_sheet("Texto OCR")
    ws.append(["ID", "Archivo", "Texto completo detectado"])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), _HEADER_FILL
    for r in records:
        ws.append([r.id, r.file_name, r.raw_text])
        ws.cell(row=ws.max_row, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 7, 28, 120
