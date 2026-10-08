"""Genera el archivo Excel con todos los cheques analizados."""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS = [
    ("N°", lambda r: r["id"]),
    ("Archivo", lambda r: r["archivo"]),
    ("Banco", lambda r: r["campos"].get("banco")),
    ("N° Cheque", lambda r: r["campos"].get("numero_cheque")),
    ("Cuenta", lambda r: r["campos"].get("cuenta")),
    ("Ciudad", lambda r: r["campos"].get("ciudad")),
    ("Fecha", lambda r: r["campos"].get("fecha")),
    ("Beneficiario", lambda r: r["campos"].get("beneficiario")),
    ("Monto", lambda r: Decimal(r.get("campos", {}).get("monto")) if r.get("campos", {}).get("monto") else None),
    ("Monto en letras", lambda r: r["campos"].get("monto_letras")),
    ("Validación monto", lambda r: r["campos"].get("validacion_monto")),
    ("Concepto", lambda r: r["campos"].get("concepto")),
    ("Línea MICR", lambda r: r["campos"].get("linea_micr")),
    ("Confianza OCR %", lambda r: round(r.get("confianza_general", 0) * 100, 1)),
    ("Requiere revisión", lambda r: "SÍ" if r.get("requiere_revision") else "NO"),
    ("Avisos", lambda r: "; ".join(r.get("avisos") or [])),
    ("Calidad imagen", lambda r: ", ".join((r.get("calidad") or {}).get("problemas") or []) or "OK"),
    ("Estado", lambda r: r.get("estado")),
    ("Tiempo (ms)", lambda r: r.get("tiempo_ms")),
    ("Fecha proceso", lambda r: r.get("fecha_proceso")),
]

COLOR_ENCABEZADO = PatternFill("solid", fgColor="1F4E78")
COLOR_ALERTA = PatternFill("solid", fgColor="F8D7DA")
COLOR_OK = PatternFill("solid", fgColor="D4EDDA")


def generar_excel(resultados: list[dict]) -> BytesIO:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Cheques"

    hoja.append([titulo for titulo, _ in COLUMNAS])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = COLOR_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center")

    col_monto = [t for t, _ in COLUMNAS].index("Monto") + 1
    col_validacion = [t for t, _ in COLUMNAS].index("Validación monto") + 1
    for registro in resultados:
        hoja.append([obtener(registro) for _, obtener in COLUMNAS])
        fila = hoja.max_row
        hoja.cell(fila, col_monto).number_format = "#,##0.00"
        validacion = hoja.cell(fila, col_validacion)
        if validacion.value == "COINCIDE":
            validacion.fill = COLOR_OK
        elif validacion.value == "NO COINCIDE":
            validacion.fill = COLOR_ALERTA

    for indice, columna in enumerate(hoja.columns, start=1):
        ancho = max(len(str(c.value or "")) for c in columna)
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 2, 10), 50)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    # Segunda hoja con el texto OCR completo, útil para auditoría
    hoja_texto = libro.create_sheet("Texto OCR")
    hoja_texto.append(["N°", "Archivo", "Texto completo"])
    for celda in hoja_texto[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = COLOR_ENCABEZADO
    for registro in resultados:
        hoja_texto.append([registro["id"], registro["archivo"], registro.get("texto_completo", "")])
        hoja_texto.cell(hoja_texto.max_row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    hoja_texto.column_dimensions["B"].width = 30
    hoja_texto.column_dimensions["C"].width = 100

    salida = BytesIO()
    libro.save(salida)
    salida.seek(0)
    return salida
