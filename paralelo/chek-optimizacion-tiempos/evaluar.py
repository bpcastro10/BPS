"""Mide acierto por campo y tiempo por cheque contra las respuestas correctas (verdad.csv).

Uso:
    .venv\\Scripts\\python evaluar.py                      (muestras\\ con muestras\\verdad.csv)
    .venv\\Scripts\\python evaluar.py carpeta --modelo v5_latin

Para cada campo cuenta 4 resultados:
    ok_auto        correcto y aceptado solo            (lo que se ahorra)
    ok_revisar     correcto pero marcado para revisar  (costo: una revisión)
    error_revisar  incorrecto y marcado para revisar   (el sistema avisó: no hay daño)
    error_auto     incorrecto y aceptado solo          (¡el peor caso! debe tender a 0)
"""

import argparse
import csv
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from poc.analizador import Analizador
from poc.comparar import CAMPOS, normalizar_valor as _norm

RAIZ = Path(__file__).parent


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser()
    p.add_argument("carpeta", nargs="?", default=str(RAIZ / "muestras"))
    p.add_argument("--dispositivo", default="cpu", help="cpu (por omisión) o cuda:0, cuda:1...")
    p.add_argument("--modelo")
    p.add_argument("--beneficiario", action="store_true")
    a = p.parse_args()
    carpeta = Path(a.carpeta)
    verdad = {r["archivo"]: r for r in csv.DictReader(open(carpeta / "verdad.csv", encoding="utf-8"))}

    config = json.loads((RAIZ / "config.json").read_text(encoding="utf-8"))
    if a.modelo:
        config["modelo_lectura"] = a.modelo
    if a.dispositivo != "cpu":
        from poc.hardware import detectar
        from poc.paralelo import config_del_proceso
        config = config_del_proceso(config, a.dispositivo, None, config.get("ejecucion", {}).get("precision_gpu", "float32"),
                                    detectar().onnx_cuda)
    an = Analizador(config, RAIZ)
    print(f"modelo {an.lector.modelo} cargado en {an.segundos_carga} s\n")

    filas, conteo = [], {c: {"ok_auto": 0, "ok_revisar": 0, "error_revisar": 0, "error_auto": 0} for c in CAMPOS}
    tiempos = []
    for archivo, v in verdad.items():
        r = an.analizar(carpeta / archivo, {"beneficiario": True} if a.beneficiario else None)
        tiempos.append(r["tiempos_s"]["total"])
        fila = {"archivo": archivo, "tiempo_s": r["tiempos_s"]["total"], "estado": r["estado"]}
        print(f"{archivo}  {r['tiempos_s']['total']:.2f} s  {r['estado']}")
        for c in CAMPOS:
            if c not in r["detalle"] or not v.get(c):
                continue
            det = r["detalle"][c]
            correcto = _norm(c, det["valor"]) == _norm(c, v[c])
            tipo = ("ok" if correcto else "error") + ("_revisar" if det["revisar"] else "_auto")
            conteo[c][tipo] += 1
            fila[c] = det["valor"]
            fila[c + "_esperado"] = v[c]
            fila[c + "_resultado"] = tipo
            marca = {"ok_auto": "✔", "ok_revisar": "✔?", "error_revisar": "✘?", "error_auto": "✘!!"}[tipo]
            print(f"   {marca:4s} {c:14s} {str(det['valor']):38s} esperado {v[c]}")
        filas.append(fila)

    print("\nResumen por campo (ok_auto / ok_revisar / error_revisar / error_auto):")
    for c, k in conteo.items():
        n = sum(k.values())
        if n:
            print(f"   {c:14s} {k['ok_auto']}/{k['ok_revisar']}/{k['error_revisar']}/{k['error_auto']}   "
                  f"acierto {100 * (k['ok_auto'] + k['ok_revisar']) / n:.0f} %   automático {100 * k['ok_auto'] / n:.0f} %")
    print(f"\nTiempo por cheque: promedio {sum(tiempos) / len(tiempos):.2f} s, máximo {max(tiempos):.2f} s "
          f"(presupuesto {config['presupuesto_segundos_por_lado']} s)")

    wb = Workbook()
    ws = wb.active
    ws.title = "Cheques"
    cols = ["archivo", "tiempo_s", "estado"] + [x for c in CAMPOS for x in (c, c + "_esperado", c + "_resultado")]
    ws.append(cols)
    colores = {"ok_auto": "C6EFCE", "ok_revisar": "FFEB9C", "error_revisar": "F8CBAD", "error_auto": "FF0000"}
    for f in filas:
        ws.append([f.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            if c.endswith("_resultado") and f.get(c):
                ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=colores[f[c]])
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws2 = wb.create_sheet("Resumen")
    ws2.append(["campo", "ok_auto", "ok_revisar", "error_revisar", "error_auto"])
    for c, k in conteo.items():
        ws2.append([c, k["ok_auto"], k["ok_revisar"], k["error_revisar"], k["error_auto"]])
    salida = RAIZ / "resultados" / f"evaluacion_{an.lector.modelo}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    salida.parent.mkdir(exist_ok=True)
    wb.save(salida)
    print(f"Excel: {salida}")


if __name__ == "__main__":
    main()
