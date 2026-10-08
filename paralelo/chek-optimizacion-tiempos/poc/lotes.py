"""Análisis de todas las imágenes de una carpeta, en segundo plano, con avance consultable.

Los lotes se procesan de a uno y comparten el motor con el análisis individual (poc/paralelo.py).
Con 1 proceso se analiza un cheque a la vez (como siempre); con N procesos, N cheques del lote a la vez.
Los resultados quedan en memoria mientras el servidor esté encendido. Al terminar se guarda un Excel
en resultados/ (en el orden de la carpeta).
"""

import csv
import threading
import time
import uuid
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from poc.comparar import CAMPOS, clasificar

EXTENSIONES = {".tif", ".tiff", ".png", ".jpg", ".jpeg"}
COLORES = {"ok_auto": "C6EFCE", "ok_revisar": "FFEB9C", "error_revisar": "F8CBAD", "error_auto": "FF0000"}


def listar_imagenes(carpeta, subcarpetas):
    patron = carpeta.rglob("*") if subcarpetas else carpeta.glob("*")
    return sorted(p for p in patron if p.is_file() and p.suffix.lower() in EXTENSIONES)


def _leer_verdad(carpeta):
    """Respuestas correctas si la carpeta trae verdad.csv (clave: nombre del archivo)."""
    ruta = carpeta / "verdad.csv"
    if not ruta.exists():
        return {}
    with open(ruta, encoding="utf-8-sig") as f:
        return {r["archivo"]: r for r in csv.DictReader(f)}


class Lote:
    def __init__(self, carpeta, imagenes, seleccion, verdad):
        self.id = uuid.uuid4().hex[:10]
        self.carpeta = carpeta
        self.imagenes = imagenes
        self.seleccion = seleccion
        self.verdad = verdad
        self.resultados = []        # resumen por imagen (sin la imagen)
        self.estado = "en_cola"     # en_cola, procesando, terminado, cancelado
        self.cancelar = False
        self.inicio = None
        self.fin = None
        self.excel = None

    def resumen(self, desde=0):
        hechos = len(self.resultados)
        tiempos = [r["tiempo_s"] for r in self.resultados if r.get("tiempo_s") is not None]
        transcurrido = (self.fin or time.time()) - self.inicio if self.inicio else 0
        restante = (transcurrido / hechos) * (len(self.imagenes) - hechos) if hechos and self.estado == "procesando" else 0
        conteo = {"OK": 0, "REVISAR": 0, "ERROR": 0}
        for r in self.resultados:
            conteo[r["estado"]] += 1
        acierto = {}
        if self.verdad:
            for c in CAMPOS:
                k = {"ok_auto": 0, "ok_revisar": 0, "error_revisar": 0, "error_auto": 0}
                for r in self.resultados:
                    t = (r.get("comparacion") or {}).get(c)
                    if t:
                        k[t] += 1
                if sum(k.values()):
                    acierto[c] = k
        return {
            "id": self.id, "carpeta": str(self.carpeta), "estado": self.estado, "total": len(self.imagenes),
            "hechos": hechos, "conteo": conteo, "con_verdad": bool(self.verdad), "acierto": acierto,
            "tiempo_promedio_s": round(sum(tiempos) / len(tiempos), 2) if tiempos else None,
            "tiempo_maximo_s": round(max(tiempos), 2) if tiempos else None,
            "segundos_transcurridos": round(transcurrido, 1), "segundos_restantes": round(restante),
            "excel": bool(self.excel),
            "resultados": self.resultados[desde:], "desde": desde,
        }


class GestorLotes:
    def __init__(self, motor, carpeta_resultados):
        self.motor = motor          # MotorLocal o PoolProcesos (poc/paralelo.py)
        self.salida = Path(carpeta_resultados)
        self.lotes = {}
        self._cola = []
        self._hay_trabajo = threading.Event()
        threading.Thread(target=self._trabajar, daemon=True).start()

    def crear(self, ruta, subcarpetas, seleccion):
        carpeta = Path(ruta.strip().strip('"'))
        if not carpeta.is_dir():
            raise ValueError(f"No existe la carpeta o no es una carpeta: {carpeta}")
        imagenes = listar_imagenes(carpeta, subcarpetas)
        if not imagenes:
            raise ValueError("La carpeta no tiene imágenes (.tif, .tiff, .png, .jpg, .jpeg)")
        lote = Lote(carpeta, imagenes, seleccion, _leer_verdad(carpeta))
        self.lotes[lote.id] = lote
        self._cola.append(lote)
        self._hay_trabajo.set()
        return lote

    def _trabajar(self):
        while True:
            self._hay_trabajo.wait()
            while self._cola:
                self._procesar(self._cola.pop(0))
            self._hay_trabajo.clear()

    def _procesar(self, lote):
        lote.estado, lote.inicio = "procesando", time.time()
        n = max(1, getattr(self.motor, "procesos", 1))
        cancelado = False
        if n == 1:   # modo original: de a un cheque, en orden
            for img in lote.imagenes:
                if lote.cancelar:
                    cancelado = True
                    break
                self._uno(lote, img)
        else:        # N procesos: N cheques a la vez (los resultados llegan según terminan)
            with ThreadPoolExecutor(max_workers=n, thread_name_prefix="lote") as ex:
                pendientes = set()
                for img in lote.imagenes:
                    if lote.cancelar:
                        cancelado = True
                        break
                    if len(pendientes) >= n:
                        _, pendientes = wait(pendientes, return_when=FIRST_COMPLETED)
                    pendientes.add(ex.submit(self._uno, lote, img))
                wait(pendientes)
        lote.estado = "cancelado" if cancelado else "terminado"
        lote.fin = time.time()
        try:
            lote.excel = self._excel(lote)
        except Exception:   # noqa: BLE001 - el lote vale aunque falle el Excel
            lote.excel = None

    def _uno(self, lote, img):
        relativo = str(img.relative_to(lote.carpeta))
        fila = {"archivo": relativo}
        try:
            r, _, _ = self.motor.analizar(img, lote.seleccion)
            fila.update(estado=r["estado"], tiempo_s=r["tiempos_s"]["total"], tiempos_s=r["tiempos_s"], campos=r["campos"],
                        detalle=r["detalle"], campos_a_revisar=r["campos_a_revisar"],
                        dentro_de_presupuesto=r["dentro_de_presupuesto"],
                        validaciones=r["validaciones"], lecturas_crudas=r["lecturas_crudas"],
                        zonas=r["localizacion"]["zonas"], localizacion_notas=r["localizacion"]["notas"])
            esperado = lote.verdad.get(img.name) or lote.verdad.get(relativo)
            if esperado:
                fila["comparacion"] = {c: t for c in CAMPOS if c in r["detalle"]
                                       if (t := clasificar(c, r["detalle"][c], esperado.get(c)))}
                fila["esperado"] = {c: esperado.get(c) for c in fila["comparacion"]}
        except Exception as e:   # una imagen mala no detiene el lote
            fila.update(estado="ERROR", tiempo_s=None, error=str(e))
        lote.resultados.append(fila)

    def _excel(self, lote):
        wb = Workbook()
        ws = wb.active
        ws.title = "Cheques"
        # Una columna de tiempo por paso, en el orden en que se ejecutan
        pasos = list(dict.fromkeys(p for f in lote.resultados for p in f.get("tiempos_s", {}) if p != "total"))
        cols = ["archivo", "estado", "tiempo_s"] + [f"t_{p}_s" for p in pasos] + ["campos_a_revisar"]
        for c in CAMPOS:
            cols += [c, f"{c}_confianza", f"{c}_revisar"] + ([f"{c}_esperado", f"{c}_resultado"] if lote.verdad else [])
        cols.append("error")
        ws.append(cols)
        # Con varios procesos los resultados llegan desordenados: el Excel sigue el orden de la carpeta
        orden = {str(p.relative_to(lote.carpeta)): i for i, p in enumerate(lote.imagenes)}
        for f in sorted(lote.resultados, key=lambda f: orden.get(f["archivo"], 0)):
            det = f.get("detalle", {})
            fila = {"archivo": f["archivo"], "estado": f["estado"], "tiempo_s": f.get("tiempo_s"),
                    "campos_a_revisar": ", ".join(f.get("campos_a_revisar", [])), "error": f.get("error")}
            for p, v in f.get("tiempos_s", {}).items():
                fila[f"t_{p}_s"] = v
            for c, d in det.items():
                fila[c], fila[f"{c}_confianza"], fila[f"{c}_revisar"] = d["valor"], d["confianza"], d["revisar"]
            for c, t in (f.get("comparacion") or {}).items():
                fila[f"{c}_resultado"], fila[f"{c}_esperado"] = t, f["esperado"][c]
            ws.append([fila.get(c) for c in cols])
            for i, c in enumerate(cols, 1):
                if c.endswith("_resultado") and fila.get(c):
                    ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=COLORES[fila[c]])
        for celda in ws[1]:
            celda.font = Font(bold=True)
        ws.freeze_panes = "B2"
        ws2 = wb.create_sheet("Resumen")
        res = lote.resumen()
        ws2.append(["Carpeta", str(lote.carpeta)])
        ws2.append(["Imágenes", res["total"]])
        for k, v in res["conteo"].items():
            ws2.append([k, v])
        ws2.append(["Tiempo promedio (s)", res["tiempo_promedio_s"]])
        ws2.append(["Tiempo máximo (s)", res["tiempo_maximo_s"]])
        if pasos:
            ws2.append([])
            ws2.append(["paso", "promedio_s", "maximo_s"])
            for p in pasos:
                vals = [f["tiempos_s"][p] for f in lote.resultados if p in f.get("tiempos_s", {})]
                ws2.append([p, round(sum(vals) / len(vals), 3), max(vals)])
        if res["acierto"]:
            ws2.append([])
            ws2.append(["campo", "ok_auto", "ok_revisar", "error_revisar", "error_auto"])
            for c, k in res["acierto"].items():
                ws2.append([c, k["ok_auto"], k["ok_revisar"], k["error_revisar"], k["error_auto"]])
        self.salida.mkdir(exist_ok=True)
        ruta = self.salida / f"lote_{lote.carpeta.name}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        wb.save(ruta)
        return ruta
