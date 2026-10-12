from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from back.dominio.comparar import CAMPOS

COLORES = {"ok_auto": "C6EFCE", "ok_revisar": "FFEB9C", "error_revisar": "F8CBAD", "error_auto": "FF0000"}

def guardar_excel_lote(lote, carpeta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cheques"
    pasos = list(dict.fromkeys(p for f in lote.resultados for p in f.get("tiempos_s", {}) if p != "total"))
    cols = ["archivo", "estado", "tiempo_s"] + [f"t_{p}_s" for p in pasos] + ["campos_a_revisar"]
    for c in CAMPOS:
        cols += [c, f"{c}_confianza", f"{c}_revisar"] + ([f"{c}_esperado", f"{c}_resultado"] if lote.verdad else [])
    cols.append("error")
    ws.append(cols)
    orden = {str(p.relative_to(lote.carpeta)): i for i, p in enumerate(lote.imagenes)}
    for f in sorted(lote.resultados, key=lambda f: orden.get(f["archivo"], 0)):
        det = f.get("detalle", {})
        fila = {"archivo": f["archivo"], "estado": f["estado"], "tiempo_s": f.get("tiempo_s"),
                "campos_a_revisar": ", ".join(f.get("campos_a_revisar", [])), "error": f.get("error")}
        for p, v in f.get("tiempos_s", {}).items():
            fila[f"t_{p}_s"] = v
        for c, d in det.items():
            fila[c], fila[f"{c}_confianza"], fila[f"{c}_revisar"] = d["valor"], d["confianza"], d["revisar"]
        for c, t in (f.get("comparacion") or {}).items():
            fila[f"{c}_resultado"], fila[f"{c}_esperado"] = t, f["esperado"][c]
        ws.append([fila.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            if c.endswith("_resultado") and fila.get(c):
                ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=COLORES[fila[c]])
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws2 = wb.create_sheet("Resumen")
    res = lote.resumen()
    ws2.append(["Carpeta", str(lote.carpeta)])
    ws2.append(["Imágenes", res["total"]])
    for k, v in res["conteo"].items():
        ws2.append([k, v])
    ws2.append(["Tiempo promedio (s)", res["tiempo_promedio_s"]])
    ws2.append(["Tiempo máximo (s)", res["tiempo_maximo_s"]])
    if pasos:
        ws2.append([])
        ws2.append(["paso", "promedio_s", "maximo_s"])
        for p in pasos:
            vals = [f["tiempos_s"][p] for f in lote.resultados if p in f.get("tiempos_s", {})]
            ws2.append([p, round(sum(vals) / len(vals), 3), max(vals)])
    if res["acierto"]:
        ws2.append([])
        ws2.append(["campo", "ok_auto", "ok_revisar", "error_revisar", "error_auto"])
        for c, k in res["acierto"].items():
            ws2.append([c, k["ok_auto"], k["ok_revisar"], k["error_revisar"], k["error_auto"]])
    carpeta = Path(carpeta)
    carpeta.mkdir(exist_ok=True)
    ruta = carpeta / f"lote_{lote.carpeta.name}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(ruta)
    return ruta

def guardar_excel_evaluacion(filas, conteo, modelo, carpeta):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cheques"
    cols = ["archivo", "tiempo_s", "estado"] + [x for c in CAMPOS for x in (c, c + "_esperado", c + "_resultado")]
    ws.append(cols)
    for f in filas:
        ws.append([f.get(c) for c in cols])
        for i, c in enumerate(cols, 1):
            if c.endswith("_resultado") and f.get(c):
                ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=COLORES[f[c]])
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws2 = wb.create_sheet("Resumen")
    ws2.append(["campo", "ok_auto", "ok_revisar", "error_revisar", "error_auto"])
    for c, k in conteo.items():
        ws2.append([c, k["ok_auto"], k["ok_revisar"], k["error_revisar"], k["error_auto"]])
    carpeta = Path(carpeta)
    carpeta.mkdir(exist_ok=True)
    ruta = carpeta / f"evaluacion_{modelo}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(ruta)
    return ruta
